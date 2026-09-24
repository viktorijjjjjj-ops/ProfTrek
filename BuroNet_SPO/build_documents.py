from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt as PptPt


ROOT = Path(__file__).resolve().parent
PINE = "174B43"
PINE_LIGHT = "22675B"
INK = "182B2B"
MUTED = "6D7D79"
PAPER = "F4F2EB"
CARD = "FFFDF8"
ACID = "D8E970"
MINT = "D9EEE1"
AMBER = "F2B45E"
CORAL = "DC725F"


def set_cell_fill(cell, color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, color=INK, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_fill(table.rows[0].cells[index], PINE)
        set_cell_text(table.rows[0].cells[index], header, True, "FFFFFF", 8)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                set_cell_fill(cells[index], "F1F3ED")
            set_cell_text(cells[index], str(value), index == 0, INK, 8.5)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        run.bold = True
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        paragraph = doc.add_paragraph(item, style=style)
        paragraph.paragraph_format.space_after = Pt(2)


def build_passport(target):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.6)
    for name, size, color in [
        ("Normal", 9.5, INK), ("Title", 24, PINE),
        ("Heading 1", 16, PINE), ("Heading 2", 12, PINE_LIGHT),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    run = title.add_run("ПАСПОРТ\nРЕГИОНАЛЬНОГО ПРОЕКТА")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(PINE_LIGHT)
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(28)
    run = name.add_run("«БЮРОНЕТ СПО»")
    run.bold = True
    run.font.name = "Georgia"
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(PINE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.add_run("Региональный цифровой агент по снижению\nбюрократической нагрузки в профессиональных\nобразовательных организациях")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(125)
    meta.add_run("Луганская Народная Республика\n2026").font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_section(WD_SECTION.NEW_PAGE)

    add_heading(doc, "1. Общие сведения")
    add_table(doc, ["Параметр", "Содержание"], [
        ("Полное наименование", "Региональный цифровой агент по снижению бюрократической нагрузки в профессиональных образовательных организациях"),
        ("Рабочее название", "«БюроНет СПО»"),
        ("Тип проекта", "Региональный цифровой организационно-управленческий проект"),
        ("Инициатор и заказчик", "Министерство образования и науки Луганской Народной Республики"),
        ("Оператор", "Региональный координационный центр СПО, определяемый заказчиком"),
        ("Участники", "ПОО, структурные подразделения Министерства, владельцы региональных информационных систем"),
        ("География", "Луганская Народная Республика"),
        ("Срок реализации", "12 месяцев до промышленного запуска с последующим постоянным развитием"),
        ("Продукт", "Защищенная платформа с агентом-помощником, единым реестром запросов и аналитикой нагрузки"),
    ], [4.4, 12.7])

    add_heading(doc, "2. Актуальность и проблема")
    add_body(doc, "Сотрудники ПОО регулярно получают запросы от разных подразделений и организаций. Одни и те же сведения повторно собираются в таблицах, письмах и формах. Значительная часть времени уходит на поиск актуальной версии, перенос данных, сверку форматов и подготовку сопроводительных документов.")
    add_body(doc, "Отсутствует единый управляемый процесс регистрации, проверки и исполнения информационных запросов. Это приводит к повторному вводу, несогласованным срокам, риску расхождений и сокращению времени на основную образовательную деятельность.")

    add_heading(doc, "3. Цель проекта")
    goal = doc.add_paragraph()
    goal.paragraph_format.space_after = Pt(7)
    goal.paragraph_format.left_indent = Cm(.5)
    goal.paragraph_format.right_indent = Cm(.5)
    run = goal.add_run("К концу первого года сократить не менее чем на 40% среднее время подготовки типового ответа ПОО и не менее чем на 25% количество повторных запросов за счет единого реестра, повторного использования проверенных данных и безопасной подготовки проектов документов цифровым агентом.")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(PINE)

    add_heading(doc, "4. Задачи")
    add_bullets(doc, [
        "Провести инвентаризацию обязательной и инициативной отчетности ПОО.",
        "Создать единый реестр запросов, сроков, оснований и ответственных.",
        "Ввести паспорт показателя: определение, единицу, период, источник и владельца.",
        "Настроить поиск полных и смысловых дублей.",
        "Подключить доверенные источники и исключить повторный ручной ввод.",
        "Автоматизировать подготовку проектов писем, таблиц и аналитических записок.",
        "Обеспечить обязательную проверку результата человеком.",
        "Сформировать аналитику нагрузки и реестр решений по ее снижению.",
    ], True)

    add_heading(doc, "5. Продукт проекта")
    add_bullets(doc, [
        "единое окно приема и регистрации запросов;",
        "нормативный календарь обязательной отчетности;",
        "реестр показателей и доверенных источников;",
        "поиск повторных и пересекающихся запросов;",
        "цифровой агент для подготовки проектов документов;",
        "маршруты проверки, согласования и подтверждения человеком;",
        "обезличенная аналитика трудозатрат, дублей и просрочек;",
        "журнал действий, версий и источников каждого значения.",
    ])

    add_heading(doc, "6. Сквозной процесс")
    add_bullets(doc, [
        "Запрос поступает в единое окно или загружается сотрудником.",
        "Агент извлекает инициатора, основание, срок, период и показатели.",
        "Показатели сопоставляются с реестром и ранее исполненными запросами.",
        "При совпадении готовится ссылка на уже переданные сведения и проект ответа.",
        "При новом запросе доступные поля заполняются из доверенных источников.",
        "Сотрудник проверяет расхождения и подтверждает документ.",
        "Система фиксирует отправку, версию, происхождение данных и трудозатраты.",
    ], True)

    add_heading(doc, "7. Роли и ответственность")
    add_table(doc, ["Роль", "Ответственность"], [
        ("Куратор проекта", "Управленческие решения, ресурсы и достижение цели"),
        ("Региональный координатор", "Реестр запросов, календарь, правила и аналитика нагрузки"),
        ("Руководитель ПОО", "Организация работы, назначение ответственных и контроль сроков"),
        ("Сотрудник ПОО", "Проверка и подтверждение подготовленного документа"),
        ("Владелец показателя", "Методика, качество, источник и актуальность данных"),
        ("Администратор", "Роли, справочники, интеграции, резервное копирование и аудит"),
        ("Эксперт по ИБ и ПДн", "Модель угроз и контроль защищенной обработки данных"),
    ], [4.4, 12.7])

    add_heading(doc, "8. Целевые показатели")
    add_table(doc, ["Показатель", "База", "Пилот", "1-й год"], [
        ("Сокращение времени подготовки ответа", "0%", "≥ 30%", "≥ 40%"),
        ("Сокращение повторных запросов", "0%", "≥ 15%", "≥ 25%"),
        ("Поля из доверенных источников", "замер", "≥ 50%", "≥ 60%"),
        ("Точность выявления дублей", "нет", "≥ 85%", "≥ 90%"),
        ("Просроченные ответы", "замер", "≤ 8%", "≤ 5%"),
        ("Документы без проверки человеком", "нет", "0%", "0%"),
        ("ПОО в цифровом контуре", "0", "5-7", "100%"),
    ], [8.2, 2.8, 2.8, 2.8])

    add_heading(doc, "9. Этапы и сроки")
    add_table(doc, ["Этап", "Срок", "Результат"], [
        ("Обследование и базовый замер", "1 месяц", "Карта отчетности, трудозатраты и владельцы данных"),
        ("Регламентирование и проектирование", "2 месяца", "Модель данных, процессы, ИБ и UX"),
        ("Разработка MVP", "3 месяца", "Реестр, календарь, поиск дублей и проекты ответов"),
        ("Подготовка пилота", "1 месяц", "Обучение, шаблоны и подключение 5-7 ПОО"),
        ("Пилотная эксплуатация", "2 месяца", "Метрики, обратная связь и корректировки"),
        ("Тиражирование", "3 месяца", "Все ПОО и приоритетные интеграции"),
    ], [6.0, 2.5, 8.4])

    add_heading(doc, "10. Ресурсы")
    add_bullets(doc, [
        "руководитель и координатор проекта;",
        "бизнес-аналитик, методолог СПО и представители пилотных ПОО;",
        "архитектор, разработчики, специалист по данным и тестировщик;",
        "специалисты по информационной безопасности и защите персональных данных;",
        "защищенная инфраструктура, резервное копирование и мониторинг.",
    ])
    add_body(doc, "Объем финансирования определяется после обследования, выбора модели размещения и уточнения состава интеграций.")

    add_heading(doc, "11. Риски и меры реагирования")
    add_table(doc, ["Риск", "Мера реагирования"], [
        ("Неполный реестр отчетности", "Инвентаризация с утверждением владельцами процессов"),
        ("Низкое качество данных", "Паспорт показателя, проверки и ответственный владелец"),
        ("Ошибочная рекомендация агента", "Ссылки на источник, порог уверенности и проверка человеком"),
        ("Сопротивление пользователей", "Совместное проектирование, обучение и полезные первые сценарии"),
        ("Риск утечки данных", "Защищенный контур, минимизация и разграничение доступа"),
        ("Автоматизация лишнего процесса", "Отмена и объединение запроса до автоматизации"),
    ], [6.0, 11.1])

    add_heading(doc, "12. Информационная безопасность")
    add_body(doc, "Агент не принимает юридически значимых решений, не подписывает и не отправляет документы самостоятельно. Доступ предоставляется по роли и организации. Все действия, версии и источники журналируются. Персональные данные обрабатываются только в защищенном контуре.")

    add_heading(doc, "13. Ожидаемый эффект")
    add_bullets(doc, [
        "высвобождение времени для образовательной и методической работы;",
        "сокращение повторного сбора и расхождений между версиями;",
        "единый прозрачный календарь и снижение просрочек;",
        "измеримая картина административной нагрузки региона;",
        "регулярная отмена и объединение избыточных запросов;",
        "повышение качества данных и скорости управленческих решений.",
    ])

    add_heading(doc, "14. Критерии успешности пилота")
    add_body(doc, "Достигнуты целевые показатели пилота, отсутствуют критические инциденты информационной безопасности, не менее 80% активных пользователей положительно оценивают полезность решения, а региональным штабом принято решение о тиражировании с утвержденным перечнем доработок.")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("БюроНет СПО  •  паспорт проекта")
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.save(target)


def build_description(source, target):
    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.6)
    for name, size, color in [
        ("Normal", 9.5, INK), ("Title", 24, PINE),
        ("Heading 1", 16, PINE), ("Heading 2", 12, PINE_LIGHT),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(85)
    run = cover.add_run("ОПИСАНИЕ\nРЕГИОНАЛЬНОГО ПРОЕКТА")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(PINE_LIGHT)
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(30)
    run = name.add_run("«БЮРОНЕТ СПО»")
    run.bold = True
    run.font.name = "Georgia"
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(PINE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(10)
    subtitle.add_run("Региональный цифровой агент по снижению\nбюрократической нагрузки в профессиональных\nобразовательных организациях")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(120)
    meta.add_run("Луганская Народная Республика\n2026").font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_section(WD_SECTION.NEW_PAGE)

    index = 1
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            index += 1
            continue
        if raw.startswith("| "):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
            if len(parsed) >= 2:
                headers = parsed[0]
                rows = [row for row in parsed[2:] if row]
                add_table(doc, headers, rows)
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)$", raw)
        if heading:
            add_heading(doc, heading.group(2), len(heading.group(1)) - 1)
        elif raw.startswith("- "):
            paragraph = doc.add_paragraph(raw[2:], style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", raw):
            paragraph = doc.add_paragraph(re.sub(r"^\d+\.\s", "", raw), style="List Number")
            paragraph.paragraph_format.space_after = Pt(2)
        else:
            add_body(doc, raw)
        index += 1

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("БюроНет СПО  •  описание проекта")
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.core_properties.title = title
    doc.save(target)


def add_shape(slide, x, y, w, h, color, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptColor.from_string(color)
    shape.line.color.rgb = PptColor.from_string(line or color)
    return shape


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False, font="Arial", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = PptPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = PptColor.from_string(color)
    return box


def add_bullet_box(slide, items, x, y, w, h, size=17, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.name = "Arial"
        paragraph.font.size = PptPt(size)
        paragraph.font.color.rgb = PptColor.from_string(color)
        paragraph.space_after = PptPt(11)
        paragraph.level = 0
        paragraph.text = "•  " + paragraph.text
    return box


def base_slide(prs, number, title, eyebrow, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PptColor.from_string(PINE if dark else PAPER)
    accent = ACID if dark else PINE_LIGHT
    title_color = "FFFFFF" if dark else INK
    add_text(slide, eyebrow.upper(), .65, .42, 8.5, .28, 9, accent, True)
    add_text(slide, f"{number:02d}", 12.0, .4, .65, .3, 10, accent, True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, .65, 1.0, 11.8, .75, 28, title_color, True, "Georgia")
    add_shape(slide, .65, 6.96, 1.0, .06, accent, False)
    add_text(slide, "БЮРОНЕТ СПО  /  РЕГИОНАЛЬНЫЙ ЦИФРОВОЙ АГЕНТ", 1.8, 6.82, 6.2, .25, 8, "A9C8BF" if dark else MUTED, True)
    return slide


def metric_card(slide, x, y, w, value, label, accent=PINE):
    add_shape(slide, x, y, w, 1.35, CARD, True, "DEDDD3")
    add_text(slide, value, x + .18, y + .18, w - .36, .45, 24, accent, True, "Georgia")
    add_text(slide, label, x + .18, y + .77, w - .36, .4, 10, MUTED)


def build_presentation(target):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = base_slide(prs, 1, "БюроНет СПО", "Региональный проект", True)
    add_text(slide, "Цифровой агент по снижению бюрократической нагрузки в профессиональных образовательных организациях", .7, 2.15, 8.5, 1.25, 23, "FFFFFF")
    add_shape(slide, 9.8, 1.7, 2.35, 2.35, ACID)
    add_text(slide, "БН", 10.15, 2.22, 1.65, .8, 41, PINE, True, "Georgia", PP_ALIGN.CENTER)
    add_text(slide, "Луганская Народная Республика  •  2026", .7, 5.82, 7, .3, 11, "B9D0C9")

    slide = base_slide(prs, 2, "Время уходит не на образование", "Проблема")
    add_bullet_box(slide, ["Одни и те же сведения запрашиваются в разных формах", "Сотрудники вручную ищут, переносят и сверяют данные", "Сроки и требования находятся в письмах и таблицах", "Регион не видит реальную стоимость каждого запроса"], .75, 2.0, 6.2, 3.7, 18)
    add_shape(slide, 7.65, 2.0, 4.65, 3.65, PINE)
    add_text(slide, "Запрос", 8.0, 2.45, 1.0, .3, 12, "FFFFFF", True)
    add_text(slide, "→", 9.05, 2.4, .45, .4, 20, ACID, True)
    add_text(slide, "Таблица", 9.5, 2.45, 1.0, .3, 12, "FFFFFF", True)
    add_text(slide, "→", 10.65, 2.4, .45, .4, 20, ACID, True)
    add_text(slide, "Письмо", 11.1, 2.45, .85, .3, 12, "FFFFFF", True)
    add_text(slide, "Повторный ввод\nРасхождения версий\nРиск просрочки", 8.05, 3.45, 3.8, 1.5, 21, ACID, True, "Georgia", PP_ALIGN.CENTER)

    slide = base_slide(prs, 3, "Цель первого года", "Измеримый результат")
    add_text(slide, "Сократить административную нагрузку без снижения качества и достоверности управленческих данных", .75, 1.9, 11.5, .9, 22, INK, True, "Georgia", PP_ALIGN.CENTER)
    metric_card(slide, 1.0, 3.25, 3.35, "−40%", "время подготовки типового ответа")
    metric_card(slide, 4.95, 3.25, 3.35, "−25%", "количество повторных запросов", CORAL)
    metric_card(slide, 8.9, 3.25, 3.35, "0%", "документов без проверки человеком", PINE_LIGHT)

    slide = base_slide(prs, 4, "Не новый отчет, а единое окно", "Решение")
    fields = [("01", "Основание"), ("02", "Срок"), ("03", "Показатели"), ("04", "Источник"), ("05", "Ответственный"), ("06", "История")]
    for index, (number, label) in enumerate(fields):
        x = .75 + (index % 3) * 4.1
        y = 2.0 + (index // 3) * 1.65
        add_shape(slide, x, y, 3.65, 1.2, CARD, True, "DEDDD3")
        add_text(slide, number, x + .2, y + .2, .55, .35, 12, PINE_LIGHT, True)
        add_text(slide, label, x + .85, y + .34, 2.5, .35, 17, INK, True)
    add_text(slide, "Каждый запрос становится управляемой задачей, а не новым письмом в общей папке.", .8, 5.6, 11.6, .45, 16, PINE, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, 5, "Как работает цифровой агент", "Сквозной процесс")
    steps = ["Запрос", "Извлечение\nтребований", "Поиск\nсовпадений", "Заполнение\nиз источников", "Проверка\nчеловеком", "Результат"]
    for index, label in enumerate(steps):
        x = .55 + index * 2.1
        color = ACID if index == 4 else (MINT if index in (2, 3) else CARD)
        add_shape(slide, x, 2.45, 1.68, 1.42, color, True, "CAD4CF")
        add_text(slide, label, x + .12, 2.84, 1.44, .65, 12, INK, True, align=PP_ALIGN.CENTER)
        if index < len(steps) - 1:
            add_text(slide, "→", x + 1.73, 2.93, .35, .3, 17, PINE_LIGHT, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Источник каждого значения и решение сотрудника сохраняются в журнале.", .8, 4.75, 11.7, .5, 17, PINE, True, "Georgia", PP_ALIGN.CENTER)

    slide = base_slide(prs, 6, "Главный сценарий: остановить дубль", "Практическая ценность")
    add_shape(slide, .75, 2.0, 3.55, 3.65, CARD, True, "DEDDD3")
    add_text(slide, "НОВЫЙ ЗАПРОС", 1.05, 2.35, 2.95, .3, 10, MUTED, True)
    add_text(slide, "Сведения о\nконтингенте", 1.05, 3.0, 2.95, 1.0, 23, INK, True, "Georgia")
    add_text(slide, "8 показателей  •  срок 24 часа", 1.05, 4.6, 2.95, .35, 11, MUTED)
    add_text(slide, "→", 4.62, 3.25, .5, .5, 26, PINE_LIGHT, True)
    add_shape(slide, 5.25, 2.0, 7.05, 3.65, PINE)
    add_text(slide, "АГЕНТ НАШЕЛ СОВПАДЕНИЕ", 5.65, 2.35, 5.8, .3, 10, ACID, True)
    add_text(slide, "100% показателей уже переданы\nв форме СПО-1 от 18 сентября", 5.65, 3.0, 5.7, 1.15, 23, "FFFFFF", True, "Georgia")
    add_text(slide, "Ссылка на источник  •  сравнение периодов  •  проект ответа", 5.65, 4.65, 5.7, .35, 11, "B9D0C9")

    slide = base_slide(prs, 7, "Что получает сотрудник ПОО", "Пользовательский эффект")
    add_bullet_box(slide, ["Единый список задач и сроков", "Готовый черновик вместо пустой формы", "Источник каждого значения", "Предупреждение о расхождениях", "Меньше повторного ввода и ручной сверки"], .9, 1.95, 5.6, 3.9, 19)
    add_shape(slide, 7.1, 1.95, 5.15, 3.9, PINE)
    add_text(slide, "36,5 часа", 7.55, 2.55, 4.2, .65, 30, ACID, True, "Georgia", PP_ALIGN.CENTER)
    add_text(slide, "высвобождено за месяц", 7.55, 3.35, 4.2, .35, 14, "FFFFFF", True, align=PP_ALIGN.CENTER)
    add_text(slide, "Время возвращается\nосновной работе", 7.55, 4.35, 4.2, .8, 18, "B9D0C9", False, "Georgia", PP_ALIGN.CENTER)

    slide = base_slide(prs, 8, "Что получает регион", "Управленческий эффект")
    cards = [("Карта нагрузки", "по ПОО и подразделениям"), ("Рейтинг запросов", "повторных и трудоемких"), ("Реестр решений", "отмена и объединение"), ("Контроль сроков", "без дополнительных таблиц")]
    for index, (title, label) in enumerate(cards):
        x = .8 + (index % 2) * 6.05
        y = 1.95 + (index // 2) * 1.75
        add_shape(slide, x, y, 5.55, 1.3, CARD if index != 2 else MINT, True, "DEDDD3")
        add_text(slide, title, x + .25, y + .24, 4.9, .35, 17, PINE, True, "Georgia")
        add_text(slide, label, x + .25, y + .76, 4.9, .3, 11, MUTED)
    add_text(slide, "Эффект измеряется в часах, исключенных операциях и сокращенных запросах.", .8, 5.7, 11.6, .35, 15, PINE, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, 9, "Безопасность по умолчанию", "Доверие и контроль")
    principles = [("Человек решает", "Агент готовит, сотрудник подтверждает"), ("Источник виден", "Каждый вывод можно проверить"), ("Контур защищен", "ПДн не уходят в публичные сервисы"), ("Действия записаны", "Версии и решения доступны для аудита")]
    for index, (title, label) in enumerate(principles):
        x = .8 + index * 3.05
        add_shape(slide, x, 2.1, 2.65, 3.35, PINE if index == 0 else CARD, True, PINE if index == 0 else "DEDDD3")
        add_text(slide, f"0{index+1}", x + .22, 2.38, .5, .3, 11, ACID if index == 0 else PINE_LIGHT, True)
        add_text(slide, title, x + .22, 3.0, 2.15, .7, 17, "FFFFFF" if index == 0 else INK, True, "Georgia")
        add_text(slide, label, x + .22, 4.05, 2.15, .75, 11, "B9D0C9" if index == 0 else MUTED)

    slide = base_slide(prs, 10, "Региональный пилот", "Проверка гипотезы")
    metric_card(slide, .8, 2.0, 2.65, "5-7", "профессиональных организаций")
    metric_card(slide, 3.75, 2.0, 2.65, "2", "наиболее трудоемких сценария", CORAL)
    metric_card(slide, 6.7, 2.0, 2.65, "8 недель", "опытной эксплуатации", PINE_LIGHT)
    metric_card(slide, 9.65, 2.0, 2.65, "2 замера", "до и после внедрения", AMBER)
    add_text(slide, "Начинаем с проверки дублей и подготовки типового ответа. Масштабируем только после подтвержденного эффекта.", 1.15, 4.35, 11.0, .9, 20, PINE, True, "Georgia", PP_ALIGN.CENTER)

    slide = base_slide(prs, 11, "Целевой эффект пилота", "Показатели")
    metrics = [("−30%", "время ответа"), ("−15%", "повторные запросы"), ("≥50%", "полей из источников"), ("≥85%", "точность поиска дублей"), ("0%", "без проверки человеком")]
    for index, (value, label) in enumerate(metrics):
        x = .55 + index * 2.52
        add_shape(slide, x, 2.2, 2.2, 2.55, PINE if index == 4 else CARD, True, "DEDDD3" if index != 4 else PINE)
        add_text(slide, value, x + .12, 2.77, 1.96, .55, 23, ACID if index == 4 else PINE, True, "Georgia", PP_ALIGN.CENTER)
        add_text(slide, label, x + .18, 3.62, 1.84, .55, 11, "FFFFFF" if index == 4 else MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Решение о тиражировании принимается по фактическим данным пилота.", .8, 5.55, 11.7, .35, 15, PINE, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, 12, "Дорожная карта", "12 месяцев")
    stages = [("1", "Обследование", "1 мес."), ("2", "Проектирование", "2 мес."), ("3", "MVP", "3 мес."), ("4", "Подготовка", "1 мес."), ("5", "Пилот", "2 мес."), ("6", "Тиражирование", "3 мес.")]
    for index, (number, title, duration) in enumerate(stages):
        x = .55 + index * 2.1
        add_shape(slide, x, 2.3, 1.68, 2.15, PINE if index == 4 else CARD, True, "DEDDD3" if index != 4 else PINE)
        add_text(slide, number, x + .16, 2.52, .4, .3, 11, ACID if index == 4 else PINE_LIGHT, True)
        add_text(slide, title, x + .16, 3.05, 1.35, .65, 13, "FFFFFF" if index == 4 else INK, True, "Georgia", PP_ALIGN.CENTER)
        add_text(slide, duration, x + .16, 3.88, 1.35, .3, 10, "B9D0C9" if index == 4 else MUTED, True, align=PP_ALIGN.CENTER)
        if index < len(stages) - 1:
            add_text(slide, "→", x + 1.72, 3.15, .34, .3, 16, PINE_LIGHT, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, 13, "Следующий шаг", "Запуск проекта", True)
    add_text(slide, "1", .85, 2.05, .45, .45, 17, ACID, True, "Georgia")
    add_text(slide, "Назначить куратора и регионального координатора", 1.45, 2.05, 10.2, .45, 20, "FFFFFF", True)
    add_text(slide, "2", .85, 3.05, .45, .45, 17, ACID, True, "Georgia")
    add_text(slide, "Выбрать 5-7 пилотных ПОО", 1.45, 3.05, 10.2, .45, 20, "FFFFFF", True)
    add_text(slide, "3", .85, 4.05, .45, .45, 17, ACID, True, "Georgia")
    add_text(slide, "Провести четырехнедельный замер фактической нагрузки", 1.45, 4.05, 10.2, .45, 20, "FFFFFF", True)
    add_shape(slide, .8, 5.25, 11.6, .85, ACID)
    add_text(slide, "Сначала измерить. Затем упростить. Только потом автоматизировать.", 1.1, 5.52, 11.0, .35, 16, PINE, True, "Georgia", PP_ALIGN.CENTER)
    prs.save(target)


def main():
    build_passport(ROOT / "Pasport_proekta_BuroNet_SPO.docx")
    build_description(
        ROOT / "Opisanie_proekta_BuroNet_SPO.md",
        ROOT / "Opisanie_proekta_BuroNet_SPO.docx",
    )
    build_presentation(ROOT / "Prezentaciya_BuroNet_SPO.pptx")
    print("Созданы паспорт, описание и презентация в", ROOT)


if __name__ == "__main__":
    main()
