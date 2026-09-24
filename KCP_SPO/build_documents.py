from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt

ROOT = Path(__file__).resolve().parent
BLUE = "17233B"
ACCENT = "2D66B3"
GOLD = "F2C25B"


def add_docx(source: Path, target: Path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(50)
    section.left_margin = Pt(60)
    section.right_margin = Pt(45)
    for style_name, size, color in [("Normal", 10, "1B2433"), ("Title", 24, BLUE), ("Heading 1", 17, BLUE), ("Heading 2", 13, ACCENT)]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    for line in source.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,2})\s+(.*)$", line)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
        elif line.startswith("|"):
            continue
        else:
            p = doc.add_paragraph(line.replace("**", ""))
            p.paragraph_format.space_after = Pt(5)
    doc.save(target)


def add_text(slide, text, x, y, w, h, size=22, color=BLUE, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Arial"
    p.font.size = PptPt(size)
    p.font.bold = bold
    p.font.color.rgb = PptColor.from_string(color)
    return box


def build_pptx(target: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slides = [
        ("КЦП СПО", "Региональная межведомственная платформа востребованности профессий и синхронизации контрольных цифр приема."),
        ("Региональная проблема", "Спрос работодателей, данные центров занятости и планы ПОО разрознены; нет единой карты дефицита по муниципалитетам."),
        ("Цель", "Связать кадровую потребность региона с объемами КЦП, возможностями ПОО и результатами трудоустройства."),
        ("Межведомственный контур", "Орган управления образованием • центры занятости • работодатели • ПОО • отраслевые органы • региональный штаб"),
        ("Единая модель данных", "Муниципалитет → отрасль → профессия → спрос → предложение ПОО → КЦП → выпуск → трудоустройство."),
        ("Индекс востребованности", "Заявки работодателей, вакансии, данные центров занятости, инвестиционные планы, трудоустройство выпускников и доступность подготовки."),
        ("Синхронизация КЦП", "Спрос → расчет дефицита → предложение ПОО → проверка мощности и лимитов → межведомственная экспертиза → утвержденный свод."),
        ("Карта дефицита", "Муниципалитеты, отрасли и профессии, где спрос не покрывается планом подготовки, а также зоны избыточного предложения."),
        ("Кабинеты участников", "Работодатель подтверждает спрос; центр занятости верифицирует рынок труда; ПОО подает предложение; регион формирует свод."),
        ("Региональный пилот", "5 муниципалитетов • 3 отрасли • 10 ПОО • 30 работодателей • 50 профессий с рассчитанной востребованностью"),
        ("Эффект", "Доказательное планирование приема, снижение избыточной подготовки, адресное закрытие кадрового дефицита и обратная связь от рынка труда."),
        ("Следующий шаг", "Утвердить региональную методику, назначить владельцев данных, подключить участников и провести первый межведомственный цикл КЦП."),
    ]
    for index, (title, text) in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = PptColor.from_string("F5F7FB" if index > 1 else BLUE)
        add_text(slide, f"КЦП СПО  /  {index:02d}", .7, .45, 4, .3, 10, GOLD, True)
        add_text(slide, title, .7, 1.55, 11.8, 1.2, 34 if index == 1 else 30, "FFFFFF" if index == 1 else BLUE, True)
        add_text(slide, text, .75, 3.15, 10.8, 1.8, 22 if index == 1 else 20, "FFFFFF" if index == 1 else "34435C")
        add_text(slide, "Платформа для доказательного формирования приема СПО", .75, 6.65, 8, .3, 10, "B5C2D6" if index == 1 else "718096")
    prs.save(target)


def main():
    add_docx(ROOT / "Opisanie_platformy_KCP_SPO.md", ROOT / "Opisanie_platformy_KCP_SPO.docx")
    add_docx(ROOT / "Pasport_proekta_KCP_SPO.md", ROOT / "Pasport_proekta_KCP_SPO.docx")
    add_docx(ROOT / "Doklad_KCP_SPO.md", ROOT / "Doklad_KCP_SPO.docx")
    add_docx(ROOT / "Iskhodnye_dannye_i_otvetstvennye_KCP_SPO.md", ROOT / "Iskhodnye_dannye_i_otvetstvennye_KCP_SPO.docx")
    add_docx(ROOT / "Opisanie_cifrovogo_pomoshchnika_KCP_SPO.md", ROOT / "Opisanie_cifrovogo_pomoshchnika_KCP_SPO.docx")
    presentation = ROOT / "Prezentaciya_KCP_SPO.pptx"
    try:
        build_pptx(presentation)
    except PermissionError:
        fallback = ROOT / "Prezentaciya_KCP_SPO_obnovlennaya.pptx"
        build_pptx(fallback)
        presentation = fallback
    print("Документы созданы в", ROOT)
    print("Презентация:", presentation.name)


if __name__ == "__main__":
    main()
