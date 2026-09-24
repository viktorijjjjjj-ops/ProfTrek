from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt as PptPt


ROOT = Path(__file__).resolve().parent
BLUE = "17233B"
ACCENT = "2D66B3"
GOLD = "F2C25B"
PALE = "E4EEFB"
INK = "1B2433"
MUTED = "758196"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=INK, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_inline_markdown(paragraph, text, size=11, color=INK):
    pieces = re.split(r"(\*\*.*?\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        bold = piece.startswith("**") and piece.endswith("**")
        content = piece[2:-2] if bold else piece
        run = paragraph.add_run(content)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)


def configure_doc(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.6)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color in [
        ("Title", 24, BLUE),
        ("Heading 1", 17, BLUE),
        ("Heading 2", 14, ACCENT),
        ("Heading 3", 12, BLUE),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "ПРОФТРЕК  |  ОТ ДИПЛОМА К ПРОФЕССИИ"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(ACCENT)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Кулешова Виктория Геннадьевна  •  ЛНР  •  2026")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title_page(document, compact=False):
    for _ in range(2 if compact else 4):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("СОЦИАЛЬНЫЙ ПРОЕКТ")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ПрофТрек")
    r.font.name = "Arial"
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ОТ ДИПЛОМА К ПРОФЕССИИ")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("Система подготовки и трудоустройства выпускников СПО\nдля закрытия кадровой потребности экономики региона")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(INK)
    for _ in range(2 if compact else 4):
        document.add_paragraph()
    details = [
        "Руководитель: Кулешова Виктория Геннадьевна",
        "Организация: [наименование организации-заявителя]",
        "При участии Базового центра карьеры Минобрнауки ЛНР и Совета по социальному партнёрству в сфере занятости молодёжи",
        "Луганская Народная Республика",
        "2026",
    ]
    for text in details:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string(MUTED if text.startswith("Организация") else INK)
    document.add_page_break()


def markdown_to_docx(source, destination, title_page=True):
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_doc(document)
    if title_page:
        add_title_page(document, compact="Pasport" in source.name)

    index = 0
    title_seen = False
    while index < len(lines):
        line = lines[index].rstrip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                width = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=width)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index in range(width):
                        text = row[col_index] if col_index < len(row) else ""
                        set_cell_text(table.cell(row_index, col_index), text.replace("**", ""), row_index == 0, "FFFFFF" if row_index == 0 else INK)
                        if row_index == 0:
                            set_cell_fill(table.cell(row_index, col_index), BLUE)
                        elif row_index % 2 == 0:
                            set_cell_fill(table.cell(row_index, col_index), "F3F6FA")
                set_repeat_table_header(table.rows[0])
                document.add_paragraph()
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if title_page and level <= 2 and not title_seen:
                title_seen = True
                index += 1
                continue
            document.add_heading(text, level=min(level, 3))
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)
        if numbered:
            p = document.add_paragraph(style="List Number")
            add_inline_markdown(p, numbered.group(2))
        elif bullet:
            p = document.add_paragraph(style="List Bullet")
            add_inline_markdown(p, bullet.group(1))
        else:
            p = document.add_paragraph()
            if line.startswith("**") and line.endswith("**"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(p, line.replace("  ", ""))
        index += 1

    document.save(destination)


def ppt_add_text(slide, text, left, top, width, height, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Arial"
    paragraph.font.size = PptPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = PptColor.from_string(color)
    return box


def ppt_rect(slide, left, top, width, height, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptColor.from_string(fill)
    shape.line.color.rgb = PptColor.from_string(line or fill)
    return shape


def add_slide_base(prs, number, eyebrow, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PptColor.from_string("F5F7FB")
    ppt_rect(slide, 0, 0, 0.18, 7.5, GOLD, False)
    ppt_add_text(slide, eyebrow.upper(), 0.65, 0.46, 8.8, 0.3, 10, ACCENT, True)
    ppt_add_text(slide, title, 0.65, 0.82, 11.7, 0.7, 27, BLUE, True)
    if subtitle:
        ppt_add_text(slide, subtitle, 0.67, 1.48, 11.5, 0.45, 12, MUTED)
    ppt_add_text(slide, f"{number:02d}", 12.05, 0.45, 0.62, 0.3, 10, MUTED, True, PP_ALIGN.RIGHT)
    ppt_add_text(slide, "ПРОФТРЕК", 10.9, 7.08, 1.7, 0.22, 8, ACCENT, True, PP_ALIGN.RIGHT)
    return slide


def add_bullets(slide, items, left=0.75, top=2.05, width=11.5, size=18, gap=0.72):
    for idx, item in enumerate(items):
        y = top + idx * gap
        ppt_rect(slide, left, y + 0.08, 0.12, 0.12, GOLD, True)
        ppt_add_text(slide, item, left + 0.3, y, width - 0.3, 0.52, size, INK)


def add_metric(slide, left, top, value, label, color=ACCENT):
    ppt_rect(slide, left, top, 2.75, 1.35, "FFFFFF", True, "E1E6EF")
    ppt_add_text(slide, value, left + 0.18, top + 0.15, 2.35, 0.5, 27, color, True)
    ppt_add_text(slide, label, left + 0.18, top + 0.78, 2.35, 0.35, 11, MUTED)


def build_presentation(destination):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PptColor.from_string(BLUE)
    ppt_rect(slide, 0.72, 0.68, 0.75, 0.75, GOLD, True)
    ppt_add_text(slide, "PT", 0.84, 0.89, 0.5, 0.24, 14, BLUE, True, PP_ALIGN.CENTER)
    ppt_add_text(slide, "СОЦИАЛЬНЫЙ ПРОЕКТ", 1.67, 0.77, 3.2, 0.32, 11, GOLD, True)
    ppt_add_text(slide, "ПрофТрек", 0.75, 2.0, 7.5, 0.88, 38, "FFFFFF", True)
    ppt_add_text(slide, "От диплома к профессии", 0.78, 2.9, 7.5, 0.5, 21, GOLD, True)
    ppt_add_text(slide, "Выпускники СПО для кадровой потребности\nэкономики Луганской Народной Республики", 0.8, 3.65, 7.6, 1.0, 19, "FFFFFF")
    ppt_add_text(slide, "Кулешова Виктория Геннадьевна\nПри участии Базового центра карьеры Минобрнауки ЛНР и Совета по социальному партнёрству\nЛуганская Народная Республика • 2026", 0.8, 5.48, 7.8, 1.05, 11, "BBC7D9")
    ppt_rect(slide, 9.2, 1.6, 3.25, 4.45, ACCENT, True)
    ppt_add_text(slide, "100", 9.65, 2.12, 2.2, 0.65, 34, "FFFFFF", True, PP_ALIGN.CENTER)
    ppt_add_text(slide, "участников", 9.65, 2.8, 2.2, 0.35, 14, PALE, False, PP_ALIGN.CENTER)
    ppt_add_text(slide, "70+", 9.65, 3.55, 2.2, 0.65, 34, GOLD, True, PP_ALIGN.CENTER)
    ppt_add_text(slide, "трудоустроенных", 9.45, 4.22, 2.6, 0.35, 14, PALE, False, PP_ALIGN.CENTER)
    ppt_add_text(slide, "12 месяцев", 9.5, 5.02, 2.5, 0.35, 18, "FFFFFF", True, PP_ALIGN.CENTER)
    ppt_add_text(slide, "сопровождения", 9.5, 5.42, 2.5, 0.3, 12, PALE, False, PP_ALIGN.CENTER)

    slide = add_slide_base(prs, 2, "Актуальность", "Экономике региона нужны квалифицированные кадры", "Необходимо связать профессию выпускника с реальной кадровой заявкой предприятия")
    add_bullets(slide, [
        "Предприятия ЛНР сообщают о потребности в молодых квалифицированных работниках",
        "Выпускник не всегда выходит на место, соответствующее его профессии или специальности",
        "Кадровые заявки, подготовка и трудоустройство пока недостаточно связаны",
        "Без наставника возрастает риск раннего увольнения и ухода из профессии",
        "Региону нужен единый механизм подбора и закрепления выпускников СПО",
    ], top=2.15, size=17, gap=0.78)

    slide = add_slide_base(prs, 3, "Цель проекта", "Закрыть кадровую потребность выпускниками СПО", "Кадровая заявка → подготовленный выпускник → предприятие → закрепление")
    ppt_rect(slide, 0.75, 2.05, 11.8, 1.35, BLUE, True)
    ppt_add_text(slide, "Создать систему адресной подготовки и трудоустройства 100 выпускников СПО на востребованные рабочие места предприятий ЛНР", 1.08, 2.34, 11.1, 0.8, 21, "FFFFFF", True, PP_ALIGN.CENTER)
    add_metric(slide, 0.75, 4.15, "100", "индивидуальных маршрутов")
    add_metric(slide, 3.72, 4.15, "70+", "трудоустроенных", GOLD)
    add_metric(slide, 6.69, 4.15, "20+", "работодателей-партнёров")
    add_metric(slide, 9.66, 4.15, "12", "месяцев мониторинга", GOLD)

    slide = add_slide_base(prs, 4, "Целевая аудитория", "Исключительно выпускники СПО", "Студент входит в проект на выпускном курсе как будущий выпускник")
    cards = [
        ("Профессия", "Выпускник получил конкретную квалификацию в техникуме или колледже"),
        ("Кадровый спрос", "Профессия сопоставляется с реальной потребностью предприятий региона"),
        ("Рабочее место", "Результат — трудоустройство по профилю и устойчивое закрепление"),
    ]
    for idx, (heading, body) in enumerate(cards):
        x = 0.75 + idx * 4.0
        ppt_rect(slide, x, 2.15, 3.68, 3.55, "FFFFFF", True, "E1E6EF")
        ppt_rect(slide, x + 0.2, 2.4, 0.48, 0.48, GOLD if idx == 1 else PALE, True)
        ppt_add_text(slide, str(idx + 1), x + 0.29, 2.51, 0.3, 0.18, 11, BLUE, True, PP_ALIGN.CENTER)
        ppt_add_text(slide, heading, x + 0.25, 3.12, 3.1, 0.4, 19, BLUE, True)
        ppt_add_text(slide, body, x + 0.25, 3.75, 3.12, 1.3, 14, INK)

    slide = add_slide_base(prs, 5, "Механизм", "Восемь шагов профессионального маршрута")
    steps = ["Регистрация", "Диагностика", "Маршрут", "Подготовка", "Вакансия", "Работа", "Адаптация", "Закрепление"]
    for idx, step in enumerate(steps):
        row, col = divmod(idx, 4)
        x = 0.75 + col * 3.0
        y = 2.0 + row * 1.85
        ppt_rect(slide, x, y, 2.68, 1.28, "FFFFFF", True, "E1E6EF")
        ppt_add_text(slide, f"{idx + 1:02d}", x + 0.18, y + 0.17, 0.5, 0.3, 11, ACCENT, True)
        ppt_add_text(slide, step, x + 0.18, y + 0.57, 2.25, 0.35, 16, BLUE, True)

    slide = add_slide_base(prs, 6, "Подготовка", "Навыки для уверенного выхода на рынок труда")
    add_bullets(slide, [
        "Диагностика профессиональных интересов и факторов риска",
        "Индивидуальный план действий и подбор подходящих предприятий",
        "Резюме, сопроводительное письмо и проверка вакансий",
        "Тренировочное собеседование и деловая коммуникация",
        "Трудовое законодательство и финансовая грамотность",
        "Ярмарка вакансий, экскурсии и профессиональные пробы",
    ], top=2.0, size=16, gap=0.69)

    slide = add_slide_base(prs, 7, "Первое рабочее место", "Адаптационная карта на 90 дней", "Куратор колледжа и наставник предприятия действуют совместно")
    phases = [
        ("1-я неделя", "Знакомство с местом, коллективом, правилами и безопасностью"),
        ("1-й месяц", "Освоение базовых операций под контролем наставника"),
        ("2-й месяц", "Самостоятельное выполнение типовых рабочих задач"),
        ("3-й месяц", "Оценка результатов и план дальнейшего развития"),
    ]
    for idx, (period, text) in enumerate(phases):
        x = 0.75 + idx * 3.0
        ppt_rect(slide, x, 2.2, 2.7, 3.35, "FFFFFF", True, "E1E6EF")
        ppt_rect(slide, x, 2.2, 2.7, 0.72, GOLD if idx == 0 else ACCENT, True)
        ppt_add_text(slide, period, x + 0.2, 2.4, 2.3, 0.3, 15, BLUE if idx == 0 else "FFFFFF", True, PP_ALIGN.CENTER)
        ppt_add_text(slide, text, x + 0.25, 3.35, 2.2, 1.45, 14, INK, False, PP_ALIGN.CENTER)

    slide = add_slide_base(prs, 8, "Партнёрство", "Кто и за что отвечает", "Центр карьеры колледжа и Базовый центр карьеры выполняют разные задачи")
    roles = [
        ("Центр карьеры ПОО", "Структурное подразделение колледжа: работает со своими студентами и выпускниками"),
        ("Базовый центр карьеры", "Региональный координатор: помогает сети, но не заменяет центры колледжей"),
        ("Совет", "Работодатели, кадровые приоритеты и снятие межведомственных барьеров"),
        ("Работодатели", "Вакансии, стажировки, трудоустройство, наставники и обратная связь"),
        ("Служба занятости", "Работает со всеми гражданами; в проекте предоставляет данные и меры поддержки"),
    ]
    for idx, (heading, text) in enumerate(roles):
        y = 1.95 + idx * 0.98
        ppt_rect(slide, 0.8, y, 2.85, 0.7, BLUE if idx < 3 else PALE, True)
        ppt_add_text(slide, heading, 1.0, y + 0.19, 2.45, 0.25, 13, "FFFFFF" if idx < 3 else BLUE, True)
        ppt_add_text(slide, text, 3.9, y + 0.12, 8.1, 0.46, 13, INK)

    slide = add_slide_base(prs, 9, "Отличие", "Специализированная система для выпускников СПО", "Центры занятости работают со всеми гражданами, «ПрофТрек» — только с выпускниками СПО")
    differences = [
        ("Обычно", "Информирование, мероприятия и фиксация факта трудоустройства"),
        ("ПрофТрек", "Профессия выпускника сопоставляется с кадровой заявкой предприятия"),
        ("Результат", "Кадровая заявка закрыта, выпускник закрепился через 3, 6 и 12 месяцев"),
        ("После выхода", "Наставник, карта адаптации на 90 дней и помощь при риске увольнения"),
        ("Ответственность", "Колледж сопровождает, Базовый центр координирует, Совет объединяет партнёров"),
    ]
    for idx, (heading, text) in enumerate(differences):
        y = 1.95 + idx * 0.98
        ppt_rect(slide, 0.8, y, 2.15, 0.7, GOLD if idx == 1 else PALE, True)
        ppt_add_text(slide, heading, 1.0, y + 0.19, 1.75, 0.25, 13, BLUE, True)
        ppt_add_text(slide, text, 3.25, y + 0.12, 8.8, 0.46, 13, INK)

    slide = add_slide_base(prs, 10, "Цифровой контур", "Реестр сопровождения «ПрофТрек»", "Данные превращаются в конкретные действия куратора")
    ppt_rect(slide, 0.78, 2.03, 7.4, 3.85, BLUE, True)
    add_bullets(slide, [
        "Карточки выпускников и уровни риска",
        "Индивидуальные планы и история сопровождения",
        "База работодателей, вакансий и стажировок",
        "Напоминания о контрольных точках",
        "Обезличенная аналитика результатов",
    ], left=1.1, top=2.45, width=6.7, size=15, gap=0.6)
    for shape in list(slide.shapes)[-15:]:
        if hasattr(shape, "text_frame"):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.color.type is not None and run.font.color.rgb == PptColor.from_string(INK):
                        run.font.color.rgb = PptColor.from_string("FFFFFF")
    add_metric(slide, 9.0, 2.05, "1 / 3 / 6 / 12", "контрольные месяцы")
    add_metric(slide, 9.0, 3.68, "3", "уровня риска", GOLD)

    slide = add_slide_base(prs, 11, "Календарь", "Основной этап и мониторинг")
    timeline = [
        ("Сен–окт", "Набор, исследование, партнёры"),
        ("Окт–ноя", "Диагностика и маршруты"),
        ("Ноя–мар", "Тренинги и интервью"),
        ("Мар–авг", "Подбор и трудоустройство"),
        ("До авг 2028", "Мониторинг закрепления"),
    ]
    for idx, (period, task) in enumerate(timeline):
        x = 0.73 + idx * 2.43
        ppt_rect(slide, x, 2.45, 2.15, 2.45, "FFFFFF", True, "E1E6EF")
        ppt_add_text(slide, period, x + 0.15, 2.78, 1.85, 0.38, 14, ACCENT, True, PP_ALIGN.CENTER)
        ppt_add_text(slide, task, x + 0.2, 3.48, 1.75, 0.9, 13, INK, False, PP_ALIGN.CENTER)
    ppt_add_text(slide, "Основной этап: сентябрь 2026 — август 2027", 0.75, 5.6, 7.0, 0.35, 14, BLUE, True)

    slide = add_slide_base(prs, 12, "Результаты", "Измеримые показатели пилота")
    add_metric(slide, 0.75, 2.0, "90+", "подготовили резюме")
    add_metric(slide, 3.72, 2.0, "80+", "прошли тренировочное интервью", GOLD)
    add_metric(slide, 6.69, 2.0, "70+", "трудоустроены")
    add_metric(slide, 9.66, 2.0, "50+", "закрывают кадровую потребность", GOLD)
    ppt_rect(slide, 0.75, 4.1, 11.65, 1.45, BLUE, True)
    ppt_add_text(slide, "≥ 70%", 1.1, 4.38, 2.0, 0.55, 28, GOLD, True)
    ppt_add_text(slide, "сохраняют занятость через 12 месяцев среди участников, достигших контрольной точки", 3.0, 4.34, 8.8, 0.7, 18, "FFFFFF", True)
    ppt_add_text(slide, "5+ центров карьеры  •  20+ работодателей  •  25+ наставников  •  100+ вакансий и стажировок", 0.85, 6.05, 11.3, 0.35, 13, ACCENT, True, PP_ALIGN.CENTER)

    slide = add_slide_base(prs, 13, "Кадровый результат", "Как закрывается потребность экономики региона")
    staffing_steps = [
        ("01", "Заявка предприятия", "Профессия, количество работников, компетенции и срок"),
        ("02", "Сопоставление", "Базовый центр и центры ПОО находят подходящих выпускников"),
        ("03", "Подготовка", "Практика, профессиональная проба и подготовка к отбору"),
        ("04", "Закрепление", "Трудоустройство, наставник и контроль 3, 6 и 12 месяцев"),
    ]
    for idx, (number, heading, text) in enumerate(staffing_steps):
        x = 0.73 + idx * 3.05
        ppt_rect(slide, x, 2.05, 2.72, 3.5, "FFFFFF", True, "E1E6EF")
        ppt_add_text(slide, number, x + 0.22, 2.32, 0.55, 0.35, 14, ACCENT, True)
        ppt_add_text(slide, heading, x + 0.22, 2.95, 2.25, 0.55, 17, BLUE, True)
        ppt_add_text(slide, text, x + 0.22, 3.8, 2.25, 1.15, 13, INK)
    ppt_add_text(slide, "Итог: выпускник работает по профессии, кадровая заявка предприятия закрыта", 0.85, 6.05, 11.5, 0.4, 15, ACCENT, True, PP_ALIGN.CENTER)

    slide = add_slide_base(prs, 14, "Устойчивость", "Что останется после завершения пилота")
    add_bullets(slide, [
        "Единый маршрут включён в постоянную работу центров карьеры ПОО",
        "Базовый центр карьеры продолжает координацию и мониторинг",
        "Совет ежегодно рассматривает результаты и кадровые приоритеты",
        "Сформированы база работодателей и сеть наставников",
        "Используются цифровой реестр, формы маршрутов и мониторинга",
        "Модель тиражируется в других организациях СПО",
    ], top=1.95, size=15, gap=0.67)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PptColor.from_string(BLUE)
    ppt_add_text(slide, "ПрофТрек", 0.8, 1.42, 11.7, 0.75, 36, "FFFFFF", True, PP_ALIGN.CENTER)
    ppt_add_text(slide, "Выпускник СПО — кадровый потенциал\nэкономики региона.", 1.35, 2.55, 10.6, 1.2, 25, GOLD, True, PP_ALIGN.CENTER)
    ppt_rect(slide, 5.55, 4.42, 2.22, 0.72, ACCENT, True)
    ppt_add_text(slide, "Спасибо за внимание", 5.65, 4.63, 2.02, 0.25, 12, "FFFFFF", True, PP_ALIGN.CENTER)
    ppt_add_text(slide, "Кулешова Виктория Геннадьевна", 3.0, 6.15, 7.3, 0.3, 12, "BBC7D9", False, PP_ALIGN.CENTER)

    prs.save(destination)


def main():
    markdown_to_docx(ROOT / "Socialny_proekt_ProfTrek.md", ROOT / "Socialny_proekt_ProfTrek_Kuleshova.docx")
    markdown_to_docx(ROOT / "Pasport_proekta_ProfTrek_Kuleshova.md", ROOT / "Pasport_proekta_ProfTrek_Kuleshova.docx")
    markdown_to_docx(ROOT / "Rech_dlya_zashchity_ProfTrek_Kuleshova.md", ROOT / "Rech_dlya_zashchity_ProfTrek_Kuleshova.docx", title_page=False)
    markdown_to_docx(ROOT / "Oprosnye_listy_ProfTrek.md", ROOT / "Oprosnye_listy_ProfTrek.docx", title_page=False)
    build_presentation(ROOT / "Prezentaciya_ProfTrek_Kuleshova.pptx")
    print("Комплект ПрофТрек сформирован")


if __name__ == "__main__":
    main()
